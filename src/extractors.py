"""
extractors.py

Pulls raw text out of a resume file, whatever format it comes in:
  - .pdf   -> text layer via pdfplumber; if a page has no text (scanned/photo),
             falls back to rendering that page as an image and OCR'ing it
  - .docx  -> paragraphs + tables via python-docx
  - .doc   -> legacy binary Word files are NOT supported directly (python-docx
             can't read them). We detect this and raise a clear error asking
             the user to re-save as .docx or .pdf.
  - images -> .png/.jpg/.jpeg/.tiff/.bmp/.webp via pytesseract OCR

Every function returns plain text (str). If nothing could be extracted,
returns an empty string rather than raising, so the pipeline can log it
as a low-quality/empty result instead of crashing the whole batch.

OCR language support:
  OCR (the scanned-PDF and image paths) uses Tesseract, which only reads
  scripts it has language data installed for. By default this asks for
  English + Nepali together ("eng+nep") since that's the mix expected here.
  Tesseract doesn't raise an error for a missing language pack - it silently
  drops it - so we check pytesseract's installed language list upfront and
  only request what's actually available, printing one warning per missing
  language rather than crashing the batch.

  To install the Nepali pack on macOS:
      brew install tesseract-lang
      tesseract --list-langs   # confirm "nep" is listed

  To add another language, just change OCR_LANGUAGES below (Tesseract
  language codes, joined with "+"), e.g. "eng+nep+hin" for Hindi too.
"""

import io
import os

import fitz  # PyMuPDF
import pdfplumber
import pytesseract
from PIL import Image
from docx import Document as DocxDocument


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp"}
MIN_CHARS_PER_PAGE_BEFORE_OCR_FALLBACK = 20  # below this, assume page is a scan
OCR_LANGUAGES = "eng+nep"  # Tesseract language codes, "+" separated

_warned_missing_langs = set()  # only print the fallback warning once per missing language


def _resolve_available_langs(requested: str) -> str:
    """Tesseract doesn't raise an error for a missing language pack - it just
    silently drops it and keeps going (confirmed by testing), so checking for
    an exception doesn't work. Instead we check pytesseract's installed
    language list upfront and only request languages that are actually there,
    warning once per missing language so it doesn't get lost in scroll."""
    requested_langs = requested.split("+")

    try:
        installed = set(pytesseract.get_languages(config=""))
    except Exception:
        # Can't even ask Tesseract what's installed - just pass the request
        # through as-is and let it do whatever it does.
        return requested

    available = [l for l in requested_langs if l in installed]
    missing = [l for l in requested_langs if l not in installed]

    for lang in missing:
        if lang not in _warned_missing_langs:
            print(
                f"[extractors] Warning: Tesseract language pack '{lang}' is not installed - "
                f"text in that script will be skipped/garbled. Install it with: "
                f"brew install tesseract-lang (then `tesseract --list-langs` to confirm)."
            )
            _warned_missing_langs.add(lang)

    return "+".join(available) if available else "eng"


def _ocr_image(img: Image.Image, lang: str = OCR_LANGUAGES) -> str:
    resolved_lang = _resolve_available_langs(lang)
    return pytesseract.image_to_string(img, lang=resolved_lang).strip()


def extract_text(path: str, lang: str = OCR_LANGUAGES) -> str:
    """Dispatch to the right extractor based on file extension.
    `lang` controls OCR language(s) for scanned PDFs/images (ignored for
    native-text PDFs and DOCX, which don't need OCR)."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(path, lang=lang)
    elif ext == ".docx":
        return _extract_docx(path)
    elif ext == ".doc":
        raise ValueError(
            "Legacy .doc files are not supported directly. "
            "Please re-save the file as .docx or export it as .pdf and re-run."
        )
    elif ext in IMAGE_EXTENSIONS:
        return _extract_image(path, lang=lang)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(path: str, lang: str = OCR_LANGUAGES) -> str:
    text_parts = []

    with pdfplumber.open(path) as pdf:
        doc_for_ocr = None  # lazily opened only if a page needs OCR

        for page_num, page in enumerate(pdf.pages):
            page_text = (page.extract_text() or "").strip()

            if len(page_text) >= MIN_CHARS_PER_PAGE_BEFORE_OCR_FALLBACK:
                text_parts.append(page_text)
                continue

            # Likely a scanned page -> render it as an image and OCR it
            if doc_for_ocr is None:
                doc_for_ocr = fitz.open(path)

            pix = doc_for_ocr[page_num].get_pixmap(dpi=300)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            text_parts.append(_ocr_image(img, lang=lang))

        if doc_for_ocr is not None:
            doc_for_ocr.close()

    return "\n".join(text_parts).strip()


def _extract_docx(path: str) -> str:
    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Tables often hold contact info / skills in resume templates
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def _extract_image(path: str, lang: str = OCR_LANGUAGES) -> str:
    img = Image.open(path)
    return _ocr_image(img, lang=lang)